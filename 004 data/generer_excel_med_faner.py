import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Liste over fylkeskartkontor
kontorer = [
    "Oslo", "Hamar", "Skien", "Kristiansand", "Stavanger",
    "Bergen", "Molde", "Trondheim", "Bodø", "Tromsø"
]

# Kolonner for tabellen
kolonner = [
    "Kommunenummer",
    "Geovekst-prosjektnummer",
    "Kartlegging 2026 (Pågående/Planlagt)", 
    "Forventet låseperiode (måneder)"
]

# Definer stiler
navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Kartverket Mørkeblå
light_navy_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
input_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Lys grønn
white_font = Font(color="FFFFFF", bold=True, size=12)
header_font = Font(bold=True, size=11)
bold_font = Font(bold=True)

thin_side = Side(style='thin', color="808080")
medium_side = Side(style='medium', color="1B365D")
border_thin = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_medium = Border(left=medium_side, right=medium_side, top=medium_side, bottom=medium_side)

center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", indent=1)

def generer_excel():
    filnavn = "004 data/Datainnsamling_TraktorvegSti_FINAL.xlsx"
    wb = Workbook()
    
    # Fjern standard-arket
    std = wb.active
    wb.remove(std)

    for navn in kontorer:
        ws = wb.create_sheet(title=navn)
        
        # 1. Tittel-banner (A1:D1) - Utvidet til 4 kolonner
        ws.merge_cells("A1:D1")
        cell_a1 = ws["A1"]
        cell_a1.value = f"DATAINNSAMLING: {navn.upper()}"
        cell_a1.fill = navy_fill
        cell_a1.font = white_font
        cell_a1.alignment = center_align
        ws.row_dimensions[1].height = 35

        # 2. Generell seksjon (A3:D3)
        ws["A3"] = " GENERELLE DATA FOR KONTORET"
        ws["A3"].font = header_font
        ws["A3"].fill = light_navy_fill
        ws.merge_cells("A3:D3")
        
        generelle_felter = [
            ("Kapasitet 2026 (totalt antall ukesverk):", "B4"),
            ("Forventet minimum tidsbruk per kommune (timer):", "B5"),
            ("Forventet maksimum tidsbruk per kommune (timer):", "B6")
        ]
        
        for i, (label, cell_ref) in enumerate(generelle_felter):
            row = 4 + i
            cell = ws.cell(row=row, column=1)
            cell.value = label
            cell.font = bold_font
            cell.alignment = left_align
            cell.border = border_thin
            
            # Legg til tynn ramme for de tomme cellene i raden også for symmetri
            ws.cell(row=row, column=2).border = border_thin
            ws.cell(row=row, column=3).border = border_thin
            ws.cell(row=row, column=4).border = border_thin

            input_cell = ws[cell_ref]
            input_cell.fill = input_fill
            input_cell.alignment = center_align

        # 3. Tabelloverskrifter (A8:D8)
        ws["A8"] = " KOMMUNESPESIFIKKE DATA"
        ws["A8"].font = header_font
        ws["A8"].fill = light_navy_fill
        ws.merge_cells("A8:D8")
        
        start_rad = 9
        for col_num, column_title in enumerate(kolonner, 1):
            cell = ws.cell(row=start_rad, column=col_num)
            cell.value = column_title
            cell.fill = navy_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = center_align
            cell.border = border_thin

        # 4. Tabellrader
        for row in range(start_rad + 1, start_rad + 41): # 40 rader
            for col in range(1, len(kolonner) + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = border_thin
                cell.fill = input_fill
                cell.alignment = center_align

        # 5. Justeringer
        ws.column_dimensions['A'].width = 30 # Kommunenummer
        ws.column_dimensions['B'].width = 30 # Geovekst-prosjektnummer
        ws.column_dimensions['C'].width = 40 # Kartlegging
        ws.column_dimensions['D'].width = 30 # Låseperiode
        
        ws.row_dimensions[start_rad].height = 25

    wb.save(filnavn)
    print(f"Excel-fil generert: {filnavn}")

if __name__ == "__main__":
    generer_excel()
