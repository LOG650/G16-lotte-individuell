"""Convert Peer_review_rapport.md to a Word document."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = Path(__file__).parent / "Peer_review_rapport.md"
DST = Path(__file__).parent / "Peer_review_rapport.docx"

INLINE_BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_ITALIC = re.compile(r"(?<!\*)\*([^*\n]+?)\*(?!\*)")
INLINE_CODE = re.compile(r"`([^`]+)`")


def add_runs(paragraph, text):
    """Parse simple inline markdown (**bold**, *italic*, `code`) into runs."""
    pattern = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*\n]+?\*(?!\*)|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def is_table_separator(line):
    return bool(re.match(r"^\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$", line))


def main():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    lines = SRC.read_text(encoding="utf-8").splitlines()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.rstrip()

        if not stripped.strip():
            i += 1
            continue

        if stripped.strip() == "---":
            p = doc.add_paragraph()
            p_format = p.paragraph_format
            p_format.space_before = Pt(6)
            p_format.space_after = Pt(6)
            run = p.add_run("―" * 30)
            run.font.color.rgb = None
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        if stripped.lstrip().startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            header = parse_table_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            hdr_cells = table.rows[0].cells
            for j, h in enumerate(header):
                hdr_cells[j].text = ""
                p = hdr_cells[j].paragraphs[0]
                add_runs(p, h)
                for run in p.runs:
                    run.bold = True
            for r_idx, row in enumerate(rows, start=1):
                row_cells = table.rows[r_idx].cells
                for j, val in enumerate(row[:len(header)]):
                    row_cells[j].text = ""
                    add_runs(row_cells[j].paragraphs[0], val)
            doc.add_paragraph()
            continue

        if re.match(r"^\s*[-*]\s+", stripped):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                text = re.sub(r"^\s*[-*]\s+", "", lines[i]).rstrip()
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, text)
                i += 1
            continue

        if re.match(r"^\s*\d+\.\s+", stripped):
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                text = re.sub(r"^\s*\d+\.\s+", "", lines[i]).rstrip()
                p = doc.add_paragraph(style="List Number")
                add_runs(p, text)
                i += 1
            continue

        if stripped.lstrip().startswith(">"):
            text = re.sub(r"^\s*>\s?", "", stripped)
            p = doc.add_paragraph(style="Intense Quote") if "Intense Quote" in [s.name for s in doc.styles] else doc.add_paragraph()
            add_runs(p, text)
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(DST)
    print(f"Wrote {DST}")


if __name__ == "__main__":
    main()
